import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE = "hospital.xlsx"
HEADERS = ["ID", "Name", "Age", "Disease", "Phone"]

# Create Excel file automatically
if not os.path.exists(FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Patients"
    ws.append(HEADERS)
    wb.save(FILE)

def read_rows():
    wb = load_workbook(FILE)
    ws = wb.active
    return list(ws.iter_rows(min_row=2, values_only=True))

def login():
    if user.get() == "admin" and password.get() == "admin123":
        login_win.destroy()
        dashboard()
    else:
        messagebox.showerror("Login Failed", "Use username: admin and password: admin123")

def dashboard():
    global root, name, age, disease, phone, tree, search
    root = tk.Tk()
    root.title("Hospital Management System")
    root.geometry("760x520")

    tk.Label(root, text="HOSPITAL MANAGEMENT SYSTEM", font=("Arial", 17, "bold")).pack(pady=8)
    form = tk.Frame(root)
    form.pack(pady=4)
    fields = ["Name", "Age", "Disease", "Phone"]
    boxes = []
    for i, label in enumerate(fields):
        tk.Label(form, text=label).grid(row=0, column=i, padx=5)
        box = tk.Entry(form, width=16)
        box.grid(row=1, column=i, padx=5, pady=4)
        boxes.append(box)
    name, age, disease, phone = boxes

    buttons = tk.Frame(root)
    buttons.pack(pady=5)
    tk.Button(buttons, text="Add", width=10, command=add_record).grid(row=0, column=0, padx=3)
    tk.Button(buttons, text="Update", width=10, command=update_record).grid(row=0, column=1, padx=3)
    tk.Button(buttons, text="Delete", width=10, command=delete_record).grid(row=0, column=2, padx=3)
    tk.Button(buttons, text="Clear", width=10, command=clear_form).grid(row=0, column=3, padx=3)
    tk.Button(buttons, text="Logout", width=10, command=root.destroy).grid(row=0, column=4, padx=3)

    search_frame = tk.Frame(root)
    search_frame.pack(pady=5)
    tk.Label(search_frame, text="Search name/disease:").pack(side="left")
    search = tk.Entry(search_frame, width=25)
    search.pack(side="left", padx=5)
    tk.Button(search_frame, text="Search", command=show_records).pack(side="left")
    tk.Button(search_frame, text="Show All", command=lambda: (search.delete(0, "end"), show_records())).pack(side="left", padx=5)

    tree = ttk.Treeview(root, columns=HEADERS, show="headings", height=13)
    for col in HEADERS:
        tree.heading(col, text=col)
        tree.column(col, width=130 if col == "Disease" else 100)
    tree.pack(fill="both", expand=True, padx=10, pady=8)
    tree.bind("<<TreeviewSelect>>", select_record)
    show_records()
    root.mainloop()

def clear_form():
    for box in (name, age, disease, phone):
        box.delete(0, "end")

def show_records():
    for item in tree.get_children():
        tree.delete(item)
    query = search.get().lower().strip()
    for row in read_rows():
        if not query or query in str(row[1]).lower() or query in str(row[3]).lower():
            tree.insert("", "end", values=row)

def add_record():
    values = [name.get().strip(), age.get().strip(), disease.get().strip(), phone.get().strip()]
    if not all(values):
        messagebox.showwarning("Missing Data", "Please fill all fields.")
        return
    try:
        if int(values[1]) <= 0:
            raise ValueError
    except ValueError:
        messagebox.showwarning("Invalid Age", "Age must be a positive number.")
        return
    wb = load_workbook(FILE)
    ws = wb.active
    new_id = max([int(r[0]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None] or [0]) + 1
    ws.append([new_id] + values)
    wb.save(FILE)
    clear_form(); show_records()
    messagebox.showinfo("Saved", "Patient record added.")

def select_record(event=None):
    selected = tree.selection()
    if not selected:
        return
    row = tree.item(selected[0], "values")
    clear_form()
    for box, value in zip((name, age, disease, phone), row[1:]):
        box.insert(0, value)

def update_record():
    selected = tree.selection()
    if not selected:
        messagebox.showwarning("Select Record", "Select a patient record first.")
        return
    values = [name.get().strip(), age.get().strip(), disease.get().strip(), phone.get().strip()]
    if not all(values):
        messagebox.showwarning("Missing Data", "Please fill all fields.")
        return
    try:
        if int(values[1]) <= 0:
            raise ValueError
    except ValueError:
        messagebox.showwarning("Invalid Age", "Age must be a positive number.")
        return
    record_id = tree.item(selected[0], "values")[0]
    wb = load_workbook(FILE); ws = wb.active
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 1).value) == str(record_id):
            for col, value in enumerate(values, 2):
                ws.cell(row, col, value)
            break
    wb.save(FILE); clear_form(); show_records()
    messagebox.showinfo("Updated", "Patient record updated.")

def delete_record():
    selected = tree.selection()
    if not selected:
        messagebox.showwarning("Select Record", "Select a patient record first.")
        return
    if not messagebox.askyesno("Confirm Delete", "Delete this patient record?"):
        return
    record_id = tree.item(selected[0], "values")[0]
    wb = load_workbook(FILE); ws = wb.active
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 1).value) == str(record_id):
            ws.delete_rows(row, 1)
            break
    wb.save(FILE); clear_form(); show_records()
    messagebox.showinfo("Deleted", "Patient record deleted.")

login_win = tk.Tk()
login_win.title("Hospital System Login")
login_win.geometry("300x210")
login_win.resizable(False, False)
tk.Label(login_win, text="ADMIN LOGIN", font=("Arial", 15, "bold")).pack(pady=12)
tk.Label(login_win, text="Username").pack()
user = tk.Entry(login_win); user.pack()
tk.Label(login_win, text="Password").pack()
password = tk.Entry(login_win, show="*"); password.pack()
tk.Button(login_win, text="Login", command=login).pack(pady=12)
login_win.mainloop()
